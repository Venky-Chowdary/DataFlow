"""Read Excel / CSV / JSON into row dicts with sheet + line provenance."""

from __future__ import annotations

import csv
import io
import json
from typing import Any

from .normalize import canonical_header

MAX_FILE_BYTES = 4 * 1024 * 1024
MAX_ROWS = 5_000
SUPPORTED = frozenset({".xlsx", ".xlsm", ".csv", ".json"})


class RuleIngestError(ValueError):
    """The file is not a rule workbook we can read."""


def _ext(filename: str) -> str:
    name = (filename or "").strip().lower()
    for suffix in (".xlsx", ".xlsm", ".xls", ".csv", ".json"):
        if name.endswith(suffix):
            return suffix
    return ""


def ingest_rule_file(filename: str, payload: bytes) -> list[dict[str, Any]]:
    """Every data row from the workbook, with ``_sheet`` and ``_row`` set."""
    if not payload:
        raise RuleIngestError("The file is empty.")
    if len(payload) > MAX_FILE_BYTES:
        raise RuleIngestError(
            f"Rule files are limited to {MAX_FILE_BYTES // (1024 * 1024)} MB."
        )
    ext = _ext(filename)
    if ext == ".xls":
        raise RuleIngestError(
            "Legacy .xls is not read. Save as .xlsx, CSV or JSON."
        )
    if ext not in SUPPORTED:
        raise RuleIngestError(
            "Upload an Excel workbook (.xlsx), a CSV, or a JSON array of rules."
        )
    if ext in {".xlsx", ".xlsm"}:
        rows = _from_xlsx(payload)
    elif ext == ".csv":
        rows = _from_csv(payload)
    else:
        rows = _from_json(payload)
    if not rows:
        raise RuleIngestError(
            "No rule rows were found. The first row must name source, "
            "destination and rule columns."
        )
    return rows[:MAX_ROWS]


def _project(headers: list[str], values: list[Any], sheet: str, row_number: int) -> dict[str, Any] | None:
    mapped: dict[str, Any] = {}
    for header, value in zip(headers, values):
        key = canonical_header(str(header or ""))
        if not key:
            continue
        text = "" if value is None else str(value).strip()
        if key == "rule":
            mapped[key] = (mapped.get(key) or "") or text
        elif text and not mapped.get(key):
            mapped[key] = text
    if not any(mapped.get(k) for k in ("source_column", "dest_column", "rule")):
        return None
    mapped["_sheet"] = sheet
    mapped["_row"] = row_number
    return mapped


def _from_csv(payload: bytes) -> list[dict[str, Any]]:
    text = payload.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    try:
        headers = next(reader)
    except StopIteration:
        return []
    out: list[dict[str, Any]] = []
    for index, raw in enumerate(reader, start=2):
        row = _project(headers, raw, "csv", index)
        if row:
            out.append(row)
    return out


def _from_json(payload: bytes) -> list[dict[str, Any]]:
    try:
        data = json.loads(payload.decode("utf-8-sig"))
    except json.JSONDecodeError as exc:
        raise RuleIngestError(f"JSON is not valid: {exc}") from exc
    if isinstance(data, dict):
        data = data.get("rules") or data.get("mappings") or data.get("rows") or []
    if not isinstance(data, list):
        raise RuleIngestError("JSON must be an array of rule objects, or {\"rules\": [...]}.")
    out: list[dict[str, Any]] = []
    for index, item in enumerate(data, start=1):
        if not isinstance(item, dict):
            continue
        headers = list(item.keys())
        values = [item.get(k) for k in headers]
        row = _project(headers, values, "json", index)
        if row:
            out.append(row)
    return out


def _from_xlsx(payload: bytes) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuleIngestError("openpyxl is required to read Excel workbooks.") from exc
    wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    out: list[dict[str, Any]] = []
    try:
        for sheet in wb.worksheets:
            rows_iter = sheet.iter_rows(values_only=True)
            try:
                header_row = next(rows_iter)
            except StopIteration:
                continue
            headers = ["" if cell is None else str(cell) for cell in header_row]
            if not any(canonical_header(h) for h in headers):
                continue
            for index, raw in enumerate(rows_iter, start=2):
                values = list(raw or ())
                row = _project(headers, values, sheet.title, index)
                if row:
                    out.append(row)
    finally:
        wb.close()
    return out
