"""Rule workbook compiler — closed forms execute, the rest stay in review."""

from __future__ import annotations

import io
import json

import pytest

from services.code_crosswalk import apply_code_crosswalk
from services.rule_compiler.classify import (
    classify_rule,
    named_rule_ref,
    named_rule_targets,
    parse_date_spec,
    parse_decode,
    parse_lookup,
    parse_lookup_spec,
    parse_sql_case,
    unknown_code_policy,
)
from services.rule_compiler.compile import compile_rule_workbook
from services.rule_compiler.ingest import RuleIngestError, ingest_rule_file
from services.rule_compiler.match import name_similarity, unique_linguistic_match
from services.rule_compiler.normalize import canonical_header, resolve_name, resolve_name_ex
from services.rule_compiler.roles import infer_header_roles


def test_direct_and_empty_are_the_same_closed_form():
    for text in ("", "Direct", "1:1", "passthrough", "as-is", "preserve"):
        assert classify_rule(text)["kind"] == "direct", text


def test_lookup_accepts_a_single_code_pair():
    assert parse_lookup("A → ACTIVE") == {"A": "ACTIVE"}
    assert parse_lookup("YYYY → YYYY-MM-DD") == {}
    pairs = parse_lookup("A → ACTIVE, I → INACTIVE, P → PENDING")
    assert pairs == {"A": "ACTIVE", "I": "INACTIVE", "P": "PENDING"}
    assert classify_rule("A=ACTIVE; I=INACTIVE")["kind"] == "lookup"
    # Clio set-valued: many source codes, one target. Not just the last token.
    assert parse_lookup("A, I, P → ACTIVE") == {"A": "ACTIVE", "I": "ACTIVE", "P": "ACTIVE"}
    assert parse_lookup("A|I → ACTIVE") == {"A": "ACTIVE", "I": "ACTIVE"}
    pairs, blank = parse_lookup_spec("A → ACTIVE, blank → N/A")
    assert pairs == {"A": "ACTIVE"}
    assert blank == "N/A"


def test_derive_salary_times_twelve():
    got = classify_rule("salary × 12")
    assert got["kind"] == "derive"
    assert "*" in got["expression"]


def test_unknown_prose_is_review_not_a_guess():
    got = classify_rule("for legacy customers use the old number unless migrated")
    assert got["kind"] == "unknown"
    assert got["plane"] == "review"
    assert classify_rule("Lowercase + validate email")["kind"] == "email"
    assert classify_rule("cast number")["kind"] == "cast_number"


def test_spoken_headers_fold_onto_compiler_columns():
    assert canonical_header("Source Column") == "source_column"
    assert canonical_header("destination_column") == "dest_column"
    assert canonical_header("Join From") == "join_from"
    assert canonical_header("Business Rule") == "rule"


def test_resolve_name_fails_closed_on_short_or_ambiguous():
    cols = ["customer_id", "order_id", "first_name"]
    assert resolve_name("first_name", cols) == "first_name"
    assert resolve_name("FIRST NAME", cols) == "first_name"
    assert resolve_name("id", cols) == ""
    assert resolve_name("customer", ["customer_id", "customer_name"]) == ""


def test_csv_compiles_customer_fixture():
    csv = (
        "Source,Source Column,Destination,Destination Column,Rule\n"
        "Customer,fname,Customer,first_name,Direct\n"
        "Customer,lname,Customer,last_name,Direct\n"
        "Customer,dob,Customer,birth_date,Convert MM/DD/YYYY → YYYY-MM-DD\n"
        "Customer,status,Customer,status,\"A → ACTIVE, I → INACTIVE\"\n"
        "Customer,salary,Customer,annual_salary,salary × 12\n"
        "Customer,email,Customer,email,Lowercase + validate email\n"
        "Customer,notes,Customer,notes,omit\n"
        "Customer,mystery,Customer,segment,use legacy id unless migrated\n"
    ).encode()
    src = ["fname", "lname", "dob", "status", "salary", "email", "notes", "mystery"]
    dst = ["first_name", "last_name", "birth_date", "status", "annual_salary", "email", "notes", "segment", "unused_flag"]
    report = compile_rule_workbook(
        "customer_rules.csv",
        csv,
        source_columns=src,
        dest_columns=dst,
        source_table="Customer",
        dest_table="Customer",
    )
    by_src = {r["source_column"]: r for r in report["rules"]}
    assert report["buckets"]["executable"] >= 6
    assert report["buckets"]["needs_confirmation"] >= 1
    assert by_src["fname"]["status"] == "executable"
    assert by_src["fname"]["dest_column"] == "first_name"
    assert by_src["status"]["code_crosswalk"]["A"] == "ACTIVE"
    assert by_src["email"]["transform"] == "email"
    assert by_src["notes"]["transform"] == "omit"
    assert by_src["dob"]["shape_step"]["op"] == "parse_date"
    assert by_src["dob"]["shape_step"]["options"]["format"] == "MM/DD/YYYY"
    assert by_src["dob"]["date_format"] == "MM/DD/YYYY"
    salary = [r for r in report["rules"] if r["kind"] == "derive"][0]
    assert salary["shape_step"]["op"] == "derive_column"
    assert salary["map_source"] == "annual_salary"
    mystery = by_src["mystery"]
    assert mystery["status"] == "needs_confirmation"
    assert mystery["kind"] == "unknown"
    assert report["unused_dest_count"] == 1
    assert "unused_flag" in report["unused_dest_columns"]
    # Unbound dest is not invented, unused dest is not written.
    assert all(r["dest_column"] != "unused_flag" or r["status"] != "executable" for r in report["rules"])


def test_json_array_is_accepted():
    payload = json.dumps([
        {"source_column": "fname", "destination_column": "first_name", "rule": "Direct"},
        {"source_column": "status", "dest_column": "status", "rule": "A → ACTIVE, I → INACTIVE"},
    ]).encode()
    report = compile_rule_workbook(
        "rules.json",
        payload,
        source_columns=["fname", "status"],
        dest_columns=["first_name", "status"],
    )
    assert report["rule_count"] == 2
    assert report["buckets"]["executable"] == 2


def test_xlsx_roundtrip():
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Rules"
    ws.append(["Source Column", "Destination Column", "Rule"])
    ws.append(["fname", "first_name", "Direct"])
    ws.append(["status", "status", "A → ACTIVE, I → INACTIVE"])
    buf = io.BytesIO()
    wb.save(buf)
    report = compile_rule_workbook(
        "rules.xlsx",
        buf.getvalue(),
        source_columns=["fname", "status"],
        dest_columns=["first_name", "status"],
    )
    assert report["buckets"]["executable"] == 2
    assert report["rules"][0]["provenance"]["sheet"] == "Rules"
    assert report["rules"][0]["provenance"]["row"] == 2


def test_unbound_source_is_not_executed():
    csv = b"Source Column,Destination Column,Rule\nbogus,first_name,Direct\n"
    report = compile_rule_workbook(
        "bad.csv",
        csv,
        source_columns=["fname"],
        dest_columns=["first_name"],
    )
    assert report["rules"][0]["status"] == "needs_confirmation"
    assert report["shape_steps"] == []


def test_concat_and_hash_and_join_are_closed_forms():
    concat = classify_rule("concat first_name and last_name")
    assert concat["kind"] == "concat"
    assert "first_name" in concat["columns"]
    assert classify_rule("hash PII")["kind"] == "hash"
    assert classify_rule("cast integer")["kind"] == "cast_integer"
    join = classify_rule("left join customers on customer_id")
    assert join["kind"] == "join"
    assert join["plane"] == "review"


def test_concat_csv_emits_shape_step():
    csv = (
        "Source Column,Destination Column,Rule\n"
        "first_name,full_name,concat first_name and last_name\n"
        "ssn,ssn_hash,hash PII\n"
        "orders,fact,left join customers on customer_id\n"
    ).encode()
    report = compile_rule_workbook(
        "more.csv",
        csv,
        source_columns=["first_name", "last_name", "ssn", "orders"],
        dest_columns=["full_name", "ssn_hash", "fact"],
    )
    by_kind = {r["kind"]: r for r in report["rules"]}
    assert by_kind["concat"]["status"] == "executable"
    assert by_kind["concat"]["shape_step"]["op"] == "concat_columns"
    assert by_kind["hash"]["transform"] == "hash_pii"
    assert by_kind["join"]["status"] == "needs_confirmation"
    assert by_kind["join"]["shape_step"] is None


def test_join_keys_on_a_direct_row_stay_in_review():
    csv = (
        "Source Column,Destination Column,Rule,Join From,Join On\n"
        "email,email,Direct,customers,customer_id\n"
    ).encode()
    report = compile_rule_workbook(
        "joins.csv",
        csv,
        source_columns=["email"],
        dest_columns=["email"],
    )
    kinds = [r["kind"] for r in report["rules"]]
    assert "direct" in kinds
    assert "join" in kinds
    join = next(r for r in report["rules"] if r["kind"] == "join")
    assert join["status"] == "needs_confirmation"


def test_empty_file_is_refused():
    with pytest.raises(RuleIngestError):
        ingest_rule_file("empty.csv", b"")


def test_legacy_xls_is_refused():
    with pytest.raises(RuleIngestError):
        ingest_rule_file("old.xls", b"not-empty")


def test_closed_form_matrix():
    cases = {
        "trim + lowercase": "case_lower",
        "title case": "title",
        "collapse whitespace": "collapse",
        "strip controls": "strip_controls",
        "hh:mm:ss": "time",
        "parse json": "json",
        "base64": "binary",
        "base64 decode": "binary",
        "NOT NULL": "contract",
        "absolute value": "absolute",
        "round to 2": "round",
        "truncate to 4": "truncate",
        "clamp 0 to 100": "clamp",
        "left(code, 3)": "substr",
        "prefix \"US-\"": "prefix",
        "suffix \"-EUR\"": "suffix",
        "default 0": "default",
        "if null then N/A": "default",
        "null if : n/a, NA": "null_if",
        "keep if amount > 0": "filter",
        "exclude rows where status = X": "filter",
        "divert rows if amount < 0": "divert",
        "required": "contract",
        "must be unique": "contract",
        "=UPPER(A2)": "case_upper",
        "=TRIM(A2)": "trim",
        "=VLOOKUP(A2,Sheet2!A:B,2,FALSE)": "join",
        "if(status = \"A\", \"ACTIVE\", \"INACTIVE\")": "derive",
        "assume timezone UTC": "timezone",
    }
    for text, kind in cases.items():
        got = classify_rule(text)
        assert got["kind"] == kind, f"{text!r} → {got['kind']} (want {kind})"


def test_qualified_source_and_title_row_csv():
    csv = (
        "Customer mapping spec\n"
        "Source Column,Destination Column,Rule\n"
        "Customer.fname,first_name,Direct\n"
        "# comment,ignored,ignored\n"
        "status,status,Direct\n"
    ).encode()
    report = compile_rule_workbook(
        "titled.csv",
        csv,
        source_columns=["fname", "status"],
        dest_columns=["first_name", "status"],
    )
    by_src = {r["source_column"]: r for r in report["rules"]}
    assert by_src["fname"]["status"] == "executable"
    assert by_src["fname"]["source_table"] == "Customer"


def test_tsv_and_semicolon_and_ndjson():
    tsv = b"Source Column\tDestination Column\tRule\nfname\tfirst_name\tDirect\n"
    report = compile_rule_workbook("rules.tsv", tsv, source_columns=["fname"], dest_columns=["first_name"])
    assert report["buckets"]["executable"] == 1

    semi = "Source Column;Destination Column;Rule\nfname;first_name;Direct\n".encode()
    report = compile_rule_workbook("rules.csv", semi, source_columns=["fname"], dest_columns=["first_name"])
    assert report["buckets"]["executable"] == 1

    ndjson = (
        '{"source_column":"fname","dest_column":"first_name","rule":"Direct"}\n'
        '{"source_column":"status","dest_column":"status","rule":"A → ACTIVE, I → INACTIVE"}\n'
    ).encode()
    report = compile_rule_workbook(
        "rules.ndjson",
        ndjson,
        source_columns=["fname", "status"],
        dest_columns=["first_name", "status"],
    )
    assert report["buckets"]["executable"] == 2


def test_lookup_sheet_pairs_and_dest_collision():
    csv = (
        "Source Column,Destination Column,From Code,To Value\n"
        "status,status,A,ACTIVE\n"
        "status,status,I,INACTIVE\n"
        "Source Column,Destination Column,Rule\n"
    ).encode()
    # The second header-looking row is data? From Code sheet only:
    csv = (
        "Source Column,From Code,To Value\n"
        "status,A,ACTIVE\n"
        "status,I,INACTIVE\n"
    ).encode()
    report = compile_rule_workbook(
        "codes.csv",
        csv,
        source_columns=["status"],
        dest_columns=["status"],
    )
    lookup = next(r for r in report["rules"] if r.get("code_crosswalk"))
    assert lookup["code_crosswalk"]["A"] == "ACTIVE"
    assert lookup["code_crosswalk"]["I"] == "INACTIVE"

    collide = (
        "Source Column,Destination Column,Rule\n"
        "fname,name,Direct\n"
        "lname,name,Direct\n"
    ).encode()
    report = compile_rule_workbook(
        "collide.csv",
        collide,
        source_columns=["fname", "lname"],
        dest_columns=["name"],
    )
    statuses = {r["source_column"]: r["status"] for r in report["rules"]}
    assert statuses["fname"] == "executable"
    assert statuses["lname"] == "conflict"


def test_same_source_two_dests_and_unmapped_source():
    csv = (
        "Source Column,Destination Column,Rule\n"
        "fname,first_name,Direct\n"
        "fname,display_name,Direct\n"
    ).encode()
    report = compile_rule_workbook(
        "fanout.csv",
        csv,
        source_columns=["fname", "extra"],
        dest_columns=["first_name", "display_name"],
    )
    dests = [r["dest_column"] for r in report["rules"] if r["status"] == "executable"]
    assert set(dests) == {"first_name", "display_name"}
    assert report["unmapped_source_count"] == 1
    assert "extra" in report["unmapped_source_columns"]


def test_filter_and_compound_trim_emit_shape_steps():
    csv = (
        "Source Column,Destination Column,Rule\n"
        "amount,amount,keep if amount > 0\n"
        "email,email,trim + lowercase\n"
        "code,code_prefix,\"left(code, 3)\"\n"
    ).encode()
    report = compile_rule_workbook(
        "shape.csv",
        csv,
        source_columns=["amount", "email", "code"],
        dest_columns=["amount", "email", "code_prefix"],
    )
    ops = [s["op"] for s in report["shape_steps"]]
    assert "filter_rows" in ops
    assert "trim" in ops
    assert "derive_column" in ops
    email = next(r for r in report["rules"] if r["source_column"] == "email")
    assert email["transform"] == "lower"
    assert email["status"] == "executable"


def test_exclude_rows_is_not_omit():
    assert classify_rule("exclude rows where status = X")["kind"] == "filter"
    assert classify_rule("omit")["kind"] == "omit"


def test_prose_still_fails_closed():
    assert classify_rule("for legacy customers use the old number unless migrated")["kind"] == "unknown"
    assert classify_rule("maybe use the other id")["kind"] == "unknown"


def test_title_row_does_not_steal_unusual_headers():
    csv = (
        "Customer mapping spec\n"
        "Orig Field,Target Name,How to convert\n"
        "fname,first_name,Direct\n"
    ).encode()
    report = compile_rule_workbook(
        "titled-unusual.csv",
        csv,
        source_columns=["fname"],
        dest_columns=["first_name"],
    )
    roles = {item["header"]: item["role"] for item in report["header_roles"]}
    assert roles["Orig Field"] == "source_column"
    assert report["rules"][0]["source_column"] == "fname"
    assert report["rules"][0]["status"] == "executable"


def test_unusual_headers_are_inferred_from_schema_not_aliases():
    """A customer sheet does not have to say 'Source Column'."""
    assert canonical_header("Orig Field") == ""
    assert canonical_header("How to convert") == ""
    csv = (
        "Orig Field,Target Name,How to convert\n"
        "fname,first_name,Direct\n"
        "status,status,trim + lowercase\n"
        "mystery,segment,use the legacy id unless migrated\n"
    ).encode()
    src = ["fname", "status", "mystery"]
    dst = ["first_name", "status", "segment"]
    report = compile_rule_workbook(
        "unusual.csv",
        csv,
        source_columns=src,
        dest_columns=dst,
    )
    roles = {item["header"]: item["role"] for item in report["header_roles"]}
    assert roles["Orig Field"] == "source_column"
    assert roles["Target Name"] == "dest_column"
    assert roles["How to convert"] == "rule"
    assert any(item["method"] in {"schema", "rule_pattern", "hint"} for item in report["header_roles"])
    by_src = {r["source_column"]: r for r in report["rules"]}
    assert by_src["fname"]["status"] == "executable"
    assert by_src["fname"]["dest_column"] == "first_name"
    assert by_src["status"]["transform"] == "lower"
    assert by_src["mystery"]["status"] == "needs_confirmation"
    assert by_src["mystery"]["kind"] == "unknown"


def test_opaque_headers_bind_from_cell_values_against_schema():
    csv = (
        "ColA,ColB,ColC\n"
        "fname,first_name,Direct\n"
        "lname,last_name,Direct\n"
        "status,status,A → ACTIVE, I → INACTIVE\n"
    ).encode()
    report = compile_rule_workbook(
        "opaque.csv",
        csv,
        source_columns=["fname", "lname", "status"],
        dest_columns=["first_name", "last_name", "status"],
    )
    roles = {item["header"]: (item["role"], item["method"]) for item in report["header_roles"]}
    assert roles["ColA"][0] == "source_column"
    assert roles["ColA"][1] == "schema"
    assert roles["ColB"][0] == "dest_column"
    assert roles["ColB"][1] == "schema"
    assert roles["ColC"][0] == "rule"
    assert report["buckets"]["executable"] == 3
    status = next(r for r in report["rules"] if r["source_column"] == "status")
    assert status["code_crosswalk"]["A"] == "ACTIVE"
    assert status["code_crosswalk"]["I"] == "INACTIVE"


def test_hint_only_headers_without_schema_stay_in_review():
    csv = (
        "Orig Field,Landing col,Instruction\n"
        "fname,first_name,Direct\n"
    ).encode()
    report = compile_rule_workbook("hints-only.csv", csv)
    assert report["header_roles"]
    assert report["rules"][0]["status"] == "needs_confirmation"
    assert any("inferred" in issue.lower() for issue in report["rules"][0]["issues"])


def test_infer_header_roles_prefers_schema_over_name_hints():
    inferred = infer_header_roles(
        ["Legacy attr", "Outbound name", "Instruction"],
        [
            {"Legacy attr": "fname", "Outbound name": "first_name", "Instruction": "Direct"},
            {"Legacy attr": "lname", "Outbound name": "last_name", "Instruction": "trim"},
        ],
        source_columns=["fname", "lname"],
        dest_columns=["first_name", "last_name"],
    )
    assert inferred.roles["Legacy attr"] == "source_column"
    assert inferred.roles["Outbound name"] == "dest_column"
    assert inferred.roles["Instruction"] == "rule"
    methods = {item["header"]: item["method"] for item in inferred.evidence}
    assert methods["Legacy attr"] == "schema"
    assert methods["Outbound name"] == "schema"


def test_linguistic_bind_is_cupid_unique_winner():
    cols = ["customer_id", "order_id", "first_name"]
    assert resolve_name("cust_id", cols) == "customer_id"
    name, method, score = resolve_name_ex("cust_id", cols)
    assert name == "customer_id"
    assert method == "linguistic"
    assert score >= 0.72
    # Ambiguous and short names stay unbound — Similarity Flooding honesty.
    assert resolve_name("customer", ["customer_id", "customer_name"]) == ""
    assert resolve_name("id", cols) == ""
    assert resolve_name("name", ["first_name", "last_name"]) == ""
    assert unique_linguistic_match("fname", ["first_name", "last_name"])[0] == ""
    assert name_similarity("customer_id", "customer_id") == 1.0


def test_structural_sql_and_nested_excel_are_closed_forms():
    assert classify_rule("CAST(amount AS INTEGER)")["kind"] == "cast_integer"
    assert classify_rule("amount::numeric")["kind"] == "cast_number"
    assert classify_rule("COALESCE(status, N/A)")["kind"] == "default"
    assert classify_rule("NULLIF(code, '')")["kind"] == "null_if"
    assert classify_rule("LEN(code)")["kind"] == "derive"
    assert classify_rule("to_date(dob)")["kind"] == "date"
    nested = classify_rule("=TRIM(UPPER(A2))")
    assert nested["kind"] == "case_upper"
    assert any(extra.get("op") == "trim" for extra in nested.get("extras") or [])
    outer = classify_rule("=UPPER(TRIM(A2))")
    assert outer["kind"] == "case_upper"
    assert any(extra.get("op") == "trim" for extra in outer.get("extras") or [])


def test_linguistic_and_sql_compile_onto_engines():
    csv = (
        "Source Column,Destination Column,Rule\n"
        "cust_id,customer_id,Direct\n"
        "amount,amount,CAST(amount AS INTEGER)\n"
        "status,status,\"COALESCE(status, N/A)\"\n"
        "Customer.First Name,first_name,=TRIM(UPPER(A2))\n"
    ).encode()
    report = compile_rule_workbook(
        "linguistic.csv",
        csv,
        source_columns=["customer_id", "amount", "status", "first_name"],
        dest_columns=["customer_id", "amount", "status", "first_name"],
    )
    by_src = {r["source_column"]: r for r in report["rules"]}
    assert by_src["customer_id"]["status"] == "executable"
    assert by_src["customer_id"]["bind_method"] == "linguistic"
    assert by_src["amount"]["transform"] == "cast_integer"
    assert by_src["status"]["shape_step"]["op"] == "default_if_null"
    assert by_src["first_name"]["transform"] == "upper"
    assert any(s["op"] == "trim" for s in report["shape_steps"])
    assert report["matcher"] == "cupid-linguistic+instance+type+constraint"
    assert report["bind_methods"].get("linguistic", 0) >= 1


def test_duplicate_headers_and_notes_sheet_are_not_silent():
    csv = (
        "Col,Col,Rule\n"
        "fname,first_name,Direct\n"
    ).encode()
    report = compile_rule_workbook(
        "dup.csv",
        csv,
        source_columns=["fname"],
        dest_columns=["first_name"],
    )
    roles = {item["header"]: item["role"] for item in report["header_roles"]}
    assert "Col" in roles and "Col_2" in roles
    assert report["buckets"]["executable"] == 1

    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    wb = Workbook()
    notes = wb.active
    notes.title = "Instructions"
    notes.append(["Read this first"])
    notes.append([
        "Please map customer fields using the Rules tab. "
        "Contact the data team for exceptions that are not listed here. "
        "Historical codes stay in the archive until legal signs off."
    ])
    rules = wb.create_sheet("Rules")
    rules.append(["Source Column", "Destination Column", "Rule"])
    rules.append(["fname", "first_name", "Direct"])
    buf = io.BytesIO()
    wb.save(buf)
    report = compile_rule_workbook(
        "mixed.xlsx",
        buf.getvalue(),
        source_columns=["fname"],
        dest_columns=["first_name"],
    )
    kinds = {item["sheet"]: item["kind"] for item in report["sheet_kinds"]}
    assert kinds.get("Instructions") == "notes"
    assert kinds.get("Rules") == "rules"
    executable = [r for r in report["rules"] if r["status"] == "executable"]
    assert len(executable) == 1
    assert executable[0]["source_column"] == "fname"
    assert any("commentary" in (r.get("rule_text") or "").lower() for r in report["rules"])


def test_date_mask_names_mm_dd_versus_dd_mm():
    md = parse_date_spec("MM/DD/YYYY → YYYY-MM-DD")
    assert md and md["format"] == "MM/DD/YYYY"
    dm = parse_date_spec("DD/MM/YYYY")
    assert dm and dm["format"] == "DD/MM/YYYY"
    ambiguous = classify_rule("MM/DD/YYYY or DD/MM/YYYY")
    assert ambiguous["kind"] == "date"
    assert ambiguous["plane"] == "review"
    bare = classify_rule("convert to date")
    assert bare["kind"] == "date"
    assert not bare.get("format")

    csv = (
        "Source Column,Destination Column,Rule\n"
        "dob,birth_date,MM/DD/YYYY → YYYY-MM-DD\n"
        "hired,hired_on,convert to date\n"
    ).encode()
    report = compile_rule_workbook(
        "dates.csv",
        csv,
        source_columns=["dob", "hired"],
        dest_columns=["birth_date", "hired_on"],
    )
    dob = next(r for r in report["rules"] if r["source_column"] == "dob")
    hired = next(r for r in report["rules"] if r["source_column"] == "hired")
    assert dob["status"] == "executable"
    assert dob["shape_step"]["op"] == "parse_date"
    assert dob["shape_step"]["options"]["format"] == "MM/DD/YYYY"
    assert hired["transform"] == "date_iso"
    assert any("MM/DD vs DD/MM" in issue for issue in hired["issues"])


def test_cleansing_chain_and_lookup_semicolon_are_different():
    chained = classify_rule("trim then lowercase then email")
    assert chained["kind"] == "email"
    ops = {extra.get("op") for extra in chained.get("extras") or []}
    assert "trim" in ops
    assert "case" in ops
    assert classify_rule("A=ACTIVE; I=INACTIVE")["kind"] == "lookup"

    csv = (
        "Source Column,Destination Column,Rule\n"
        "email,email,trim then lowercase then email\n"
    ).encode()
    report = compile_rule_workbook(
        "chain.csv",
        csv,
        source_columns=["email"],
        dest_columns=["email"],
    )
    assert report["rules"][0]["status"] == "executable"
    assert report["rules"][0]["transform"] == "email"
    assert {s["op"] for s in report["shape_steps"]} >= {"trim", "case"}


def test_orphan_enumeration_sheet_attaches_to_named_lookup():
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    wb = Workbook()
    rules = wb.active
    rules.title = "Rules"
    rules.append(["Source Column", "Destination Column", "Rule"])
    rules.append(["status", "status", "lookup"])
    codes = wb.create_sheet("Status")
    codes.append(["From Code", "To Value"])
    codes.append(["A", "ACTIVE"])
    codes.append(["I", "INACTIVE"])
    buf = io.BytesIO()
    wb.save(buf)
    report = compile_rule_workbook(
        "enums.xlsx",
        buf.getvalue(),
        source_columns=["status"],
        dest_columns=["status"],
    )
    lookup = next(r for r in report["rules"] if r.get("code_crosswalk"))
    assert lookup["status"] == "executable"
    assert lookup["code_crosswalk"]["A"] == "ACTIVE"
    assert lookup["code_crosswalk"]["I"] == "INACTIVE"
    assert any(item["pairs"] >= 2 for item in report["lookup_coverage"])

    collide = Workbook()
    ws = collide.active
    ws.title = "Rules"
    ws.append(["Source Column", "Destination Column", "Rule"])
    ws.append(["status", "status", "lookup"])
    ws.append(["state", "state", "lookup"])
    codes = collide.create_sheet("Codes")
    codes.append(["From Code", "To Value"])
    codes.append(["A", "ACTIVE"])
    buf = io.BytesIO()
    collide.save(buf)
    report = compile_rule_workbook(
        "ambiguous-enums.xlsx",
        buf.getvalue(),
        source_columns=["status", "state"],
        dest_columns=["status", "state"],
    )
    assert any("matches 2 columns" in (r.get("rule_text") or "") for r in report["rules"])
    assert not any(r.get("code_crosswalk") and r["status"] == "executable" for r in report["rules"])


def test_named_rule_ref_forms_and_bare_direct_is_not_a_catalog_hit():
    assert named_rule_ref("%EmailClean%") == "EmailClean"
    assert named_rule_ref("use EmailClean") == "EmailClean"
    assert named_rule_ref("apply EmailClean") == "EmailClean"
    assert named_rule_ref("EmailClean()") == "EmailClean"
    assert named_rule_ref("Direct") == ""
    assert named_rule_ref("use the legacy id unless migrated") == ""
    name, cols = named_rule_targets("apply EmailClean to email, alt_email")
    assert name == "EmailClean"
    assert cols == ["email", "alt_email"]


def test_unknown_code_policy_is_recorded_never_identity():
    assert unknown_code_policy("A → ACTIVE, unmapped → OTHER") == {
        "action": "default",
        "value": "OTHER",
    }
    assert unknown_code_policy("* → OTHER")["action"] == "default"
    assert unknown_code_policy("default unmapped to 0") == {"action": "default", "value": "0"}
    assert unknown_code_policy("quarantine unknown codes") == {"action": "divert", "value": ""}
    assert unknown_code_policy("A → ACTIVE") == {"action": "refuse", "value": ""}
    assert "unmapped" not in parse_lookup("A → ACTIVE, unmapped → OTHER")
    assert "*" not in parse_lookup("A → ACTIVE, * → OTHER")


def test_oracle_decode_and_sql_case_are_closed_lookups():
    decoded = parse_decode("DECODE(status, 'A', 'ACTIVE', 'I', 'INACTIVE', 'OTHER')")
    assert decoded and decoded["kind"] == "lookup"
    assert decoded["mapping"] == {"A": "ACTIVE", "I": "INACTIVE"}
    assert decoded["unmapped"] == "OTHER"
    simple = parse_sql_case(
        "CASE status WHEN 'A' THEN 'ACTIVE' WHEN 'I' THEN 'INACTIVE' ELSE 'OTHER' END"
    )
    assert simple and simple["kind"] == "lookup"
    assert simple["mapping"]["A"] == "ACTIVE"
    assert simple["unmapped"] == "OTHER"
    searched = parse_sql_case(
        "CASE WHEN status = 'A' THEN 'ACTIVE' WHEN status = 'I' THEN 'INACTIVE' END"
    )
    assert searched and searched["mapping"]["I"] == "INACTIVE"
    ranged = parse_sql_case("CASE WHEN amount > 0 THEN amount ELSE 0 END")
    assert ranged and ranged["kind"] == "derive"
    assert "if(amount > 0" in ranged["expression"]


def test_named_catalog_expands_and_missing_stays_in_review():
    csv = (
        "Rule Name,Source Column,Destination Column,Rule\n"
        "EmailClean,,,trim then lowercase then email\n"
        "TrimName,,,trim\n"
        ",email,email,%EmailClean%\n"
        ",alt_email,alt_email,use EmailClean\n"
        ",notes,notes,EmailClean()\n"
        ",fname,first_name,Direct\n"
        ",mystery,segment,%Missing%\n"
    ).encode()
    report = compile_rule_workbook(
        "named.csv",
        csv,
        source_columns=["email", "alt_email", "notes", "fname", "mystery"],
        dest_columns=["email", "alt_email", "notes", "first_name", "segment"],
    )
    assert set(report["named_rules"]) >= {"EmailClean", "TrimName"}
    by_src = {r["source_column"]: r for r in report["rules"] if r.get("source_column")}
    email = by_src["email"]
    assert email["status"] == "executable"
    assert email["named_rule"] == "EmailClean"
    assert email["transform"] == "email"
    assert email["resolved_rule"]
    assert by_src["alt_email"]["status"] == "executable"
    assert by_src["notes"]["named_rule"] == "EmailClean"
    assert by_src["fname"]["kind"] == "direct"
    assert by_src["fname"].get("named_rule") in {"", None}
    mystery = by_src["mystery"]
    assert mystery["status"] == "needs_confirmation"
    assert any("Missing" in issue for issue in mystery["issues"])


def test_nested_named_rules_cycle_and_duplicate_definitions_fail_closed():
    nested = (
        "Rule Name,Source Column,Destination Column,Rule\n"
        "Inner,,,trim\n"
        "Outer,,, %Inner% \n"
        ",email,email,%Outer%\n"
    ).encode()
    report = compile_rule_workbook(
        "nested.csv",
        nested,
        source_columns=["email"],
        dest_columns=["email"],
    )
    email = next(r for r in report["rules"] if r.get("source_column") == "email")
    assert email["status"] == "executable"
    assert email["named_rule"] == "Outer"
    assert email["transform"] == "trim"

    cycle = (
        "Rule Name,Source Column,Destination Column,Rule\n"
        "A,,,%B%\n"
        "B,,,%A%\n"
        ",email,email,%A%\n"
    ).encode()
    report = compile_rule_workbook(
        "cycle.csv",
        cycle,
        source_columns=["email"],
        dest_columns=["email"],
    )
    email = next(r for r in report["rules"] if r.get("source_column") == "email")
    assert email["status"] == "needs_confirmation"
    assert any("cycle" in issue.lower() for issue in email["issues"])

    clash = (
        "Rule Name,Source Column,Destination Column,Rule\n"
        "EmailClean,,,trim\n"
        "EmailClean,,,lowercase\n"
        ",email,email,%EmailClean%\n"
    ).encode()
    report = compile_rule_workbook(
        "clash.csv",
        clash,
        source_columns=["email"],
        dest_columns=["email"],
    )
    assert any("more than once" in (r.get("rule_text") or "") for r in report["rules"])
    email = next(r for r in report["rules"] if r.get("source_column") == "email")
    assert email["transform"] == "trim"


def test_apply_named_rule_to_many_columns():
    csv = (
        "Rule Name,Source Column,Destination Column,Rule\n"
        "EmailClean,,,lowercase + validate email\n"
        ",,,apply EmailClean to email and alt_email\n"
    ).encode()
    report = compile_rule_workbook(
        "apply-to.csv",
        csv,
        source_columns=["email", "alt_email", "phone"],
        dest_columns=["email", "alt_email", "phone"],
    )
    by_src = {r["source_column"]: r for r in report["rules"] if r.get("source_column")}
    assert by_src["email"]["status"] == "executable"
    assert by_src["alt_email"]["status"] == "executable"
    assert by_src["email"]["transform"] == "email"
    assert "phone" not in by_src or by_src.get("phone", {}).get("status") != "executable"


def test_unknown_code_policy_does_not_weaken_g20():
    csv = (
        "Source Column,Destination Column,Rule\n"
        "status,status,\"A → ACTIVE, I → INACTIVE, unmapped → OTHER\"\n"
        "state,state,\"A → ACTIVE, I → INACTIVE, quarantine unknown\"\n"
        "flag,flag,\"DECODE(flag, 'Y', 'YES', 'N', 'NO', 'OTHER')\"\n"
    ).encode()
    report = compile_rule_workbook(
        "policy.csv",
        csv,
        source_columns=["status", "state", "flag"],
        dest_columns=["status", "state", "flag"],
    )
    by_src = {r["source_column"]: r for r in report["rules"]}
    status = by_src["status"]
    assert status["status"] == "executable"
    assert status["code_crosswalk"] == {"A": "ACTIVE", "I": "INACTIVE"}
    assert "OTHER" not in status["code_crosswalk"]
    assert "*" not in status["code_crosswalk"]
    assert status["unknown_code_policy"]["action"] == "default"
    assert status["unknown_code_policy"]["value"] == "OTHER"
    assert any("G20 still refuses" in issue for issue in status["issues"])
    state = by_src["state"]
    assert state["unknown_code_policy"]["action"] == "divert"
    flag = by_src["flag"]
    assert flag["code_crosswalk"] == {"Y": "YES", "N": "NO"}
    assert flag["unknown_code_policy"]["value"] == "OTHER"
    rewritten, err = apply_code_crosswalk(
        "Z",
        {"source": "status", "target": "status", "code_crosswalk": status["code_crosswalk"]},
    )
    assert rewritten is None
    assert err and "unmapped" in err


def test_name_expression_catalog_headers_are_inferred():
    csv = (
        "Name,Expression\n"
        "EmailClean,lowercase + validate email\n"
        "TrimName,trim\n"
    ).encode()
    report = compile_rule_workbook("catalog.csv", csv)
    roles = {item["header"]: item["role"] for item in report["header_roles"]}
    assert roles.get("Name") == "rule_name"
    assert roles.get("Expression") == "rule"
    assert set(report["named_rules"]) >= {"EmailClean", "TrimName"}
    kinds = {item["kind"] for item in report["sheet_kinds"]}
    assert "catalog" in kinds


def test_mapplet_header_alias():
    assert canonical_header("Mapplet") == "rule_name"
    assert canonical_header("Macro Name") == "rule_name"


def test_quoted_csv_keeps_decode_commas_inside_the_rule_cell():
    """csv.Sniffer must not steal quotechar from Oracle single quotes."""
    csv = (
        "Source Column,Destination Column,Rule\n"
        "status,status,\"A → ACTIVE, I → INACTIVE, unmapped → OTHER\"\n"
        "flag,flag,\"DECODE(flag, 'Y', 'YES', 'N', 'NO', 'OTHER')\"\n"
    ).encode()
    report = compile_rule_workbook(
        "quoted.csv",
        csv,
        source_columns=["status", "flag"],
        dest_columns=["status", "flag"],
    )
    by_src = {r["source_column"]: r for r in report["rules"]}
    assert by_src["status"]["code_crosswalk"] == {"A": "ACTIVE", "I": "INACTIVE"}
    assert by_src["flag"]["code_crosswalk"] == {"Y": "YES", "N": "NO"}
    assert by_src["flag"]["unknown_code_policy"]["value"] == "OTHER"


def test_clio_many_to_one_and_blank_are_not_silent():
    csv = (
        "Source Column,Destination Column,Rule\n"
        "status,status,\"A, I, P → ACTIVE, blank → UNKNOWN\"\n"
        "state,state,A|I → OPEN\n"
    ).encode()
    report = compile_rule_workbook(
        "clio.csv",
        csv,
        source_columns=["status", "state"],
        dest_columns=["status", "state"],
    )
    by_src = {r["source_column"]: r for r in report["rules"]}
    assert by_src["status"]["code_crosswalk"] == {"A": "ACTIVE", "I": "ACTIVE", "P": "ACTIVE"}
    assert "" not in by_src["status"]["code_crosswalk"]
    assert "UNKNOWN" not in by_src["status"]["code_crosswalk"]
    assert any(s["op"] == "default_if_null" for s in report["shape_steps"])
    assert by_src["state"]["code_crosswalk"] == {"A": "OPEN", "I": "OPEN"}


def test_informatica_iif_nvl2_and_update_strategy():
    iif = classify_rule("IIF(amount > 0, amount, 0)")
    assert iif["kind"] == "derive"
    assert "if(amount > 0" in iif["expression"]
    nulls = classify_rule("IIF(ISNULL(status), N/A, status)")
    assert "is_null(status)" in nulls["expression"]
    nvl2 = classify_rule("NVL2(email, email, unknown)")
    assert nvl2["kind"] == "derive"
    assert "is_not_null(email)" in nvl2["expression"]
    reject = classify_rule("IIF(ISNULL(ITEM_NAME), DD_REJECT, DD_INSERT)")
    assert reject["kind"] == "unknown"
    assert reject["plane"] == "review"
    ternary = classify_rule("status = A ? ACTIVE : INACTIVE")
    assert ternary["kind"] == "derive"
    assert "if(" in ternary["expression"]


def test_in_between_filters_and_regex_and_hash_identity():
    keep_in = classify_rule("keep if status in (A, I, P)")
    assert keep_in["kind"] == "filter"
    assert 'status = "A"' in keep_in["condition"]
    assert keep_in["keep"] is True
    exclude_in = classify_rule("exclude rows where status in (X, Y)")
    assert exclude_in["kind"] == "filter"
    assert exclude_in["keep"] is False
    between = classify_rule("keep if amount between 0 and 100")
    assert between["kind"] == "filter"
    assert "amount >= 0" in between["condition"]
    extracted = classify_rule("REG_EXTRACT(name, '([A-Za-z]+)', 1)")
    assert extracted["kind"] == "derive"
    assert "regex_extract" in extracted["expression"]
    replaced = classify_rule("REG_REPLACE(code, '[^0-9]', '')")
    assert replaced["kind"] == "replace"
    assert replaced.get("regex") is True
    hashed = classify_rule("hash identity of customer_id and email")
    assert hashed["kind"] == "hash_identity"
    assert hashed["columns"] == ["customer_id", "email"]
    assert classify_rule("hash PII")["kind"] == "hash"
    assert classify_rule("fill down account_id")["kind"] == "unknown"
    assert classify_rule("unnest json")["kind"] == "unnest"

    csv = (
        "Source Column,Destination Column,Rule\n"
        "status,status,\"keep if status in (A, I)\"\n"
        "amount,amount,keep if amount between 0 and 100\n"
        "name,first_token,\"REG_EXTRACT(name, '([A-Za-z]+)', 1)\"\n"
        "customer_id,row_key,hash identity of customer_id and email\n"
        "payload,item,unnest json\n"
        "code,code,\"REG_REPLACE(code, '[^0-9]', '')\"\n"
    ).encode()
    report = compile_rule_workbook(
        "imap.csv",
        csv,
        source_columns=["status", "amount", "name", "customer_id", "email", "payload", "code"],
        dest_columns=["status", "amount", "first_token", "row_key", "item", "code"],
    )
    ops = {s["op"] for s in report["shape_steps"]}
    assert "filter_rows" in ops
    assert "hash_identity" in ops
    assert "unnest_json" in ops
    assert "replace" in ops
    hash_step = next(s for s in report["shape_steps"] if s["op"] == "hash_identity")
    assert set(hash_step["options"]["columns"]) >= {"customer_id", "email"}
    assert hash_step["options"]["to"] == "row_key"
    unnest = next(r for r in report["rules"] if r["kind"] == "unnest")
    assert unnest["status"] == "executable"
    assert any("expanded image" in issue for issue in unnest["issues"])
    replace = next(r for r in report["rules"] if r["kind"] == "replace")
    assert replace["shape_step"]["options"].get("regex") is True


def test_like_is_null_not_in_and_keep_drop_columns():
    like = classify_rule("keep if name like 'Acme%'")
    assert like["kind"] == "filter"
    assert 'starts_with(name, "Acme")' in like["condition"]
    contains = classify_rule("keep if email like '%@corp.com'")
    assert "ends_with(email" in contains["condition"]
    mid = classify_rule("keep if note like '%urgent%'")
    assert "contains(note" in mid["condition"]
    wild = classify_rule("keep if code like 'A_B'")
    assert wild["kind"] == "unknown"
    nulls = classify_rule("keep if status is null")
    assert nulls["condition"] == "is_null(status)"
    present = classify_rule("exclude rows where email is not null")
    assert present["kind"] == "filter"
    assert present["keep"] is False
    assert "is_not_null(email)" in present["condition"]
    not_in = classify_rule("keep if status not in (X, Y)")
    assert not_in["kind"] == "filter"
    assert not_in["condition"].startswith("not ")
    keep = classify_rule("keep only columns fname, lname, email")
    assert keep["kind"] == "keep_columns"
    assert keep["columns"] == ["fname", "lname", "email"]
    drop = classify_rule("drop column notes")
    assert drop["kind"] == "drop_column"
    assert classify_rule("omit")["kind"] == "omit"
    assert classify_rule("flatten json")["kind"] == "flatten"

    csv = (
        "Source Column,Destination Column,Rule\n"
        "name,name,\"keep if name like 'Acme%'\"\n"
        "status,status,keep if status is not null\n"
        ",,keep only columns fname, lname\n"
        "notes,notes,drop column notes\n"
        "payload,payload,flatten json\n"
    ).encode()
    report = compile_rule_workbook(
        "sql-pred.csv",
        csv,
        source_columns=["name", "status", "fname", "lname", "notes", "payload"],
        dest_columns=["name", "status", "notes", "payload"],
    )
    ops = {s["op"] for s in report["shape_steps"]}
    assert "filter_rows" in ops
    assert "keep_columns" in ops
    assert "drop_column" in ops
    assert "flatten_json" in ops
    flatten = next(s for s in report["shape_steps"] if s["op"] == "flatten_json")
    assert flatten["options"]["depth"] == "top"
    notes = next(r for r in report["rules"] if r.get("source_column") == "notes")
    assert notes["transform"] == "omit"


def test_coma_type_constraint_and_cdc_refuse_preload():
    csv = (
        "Source Column,Destination Column,Rule\n"
        "dob,age,MM/DD/YYYY → YYYY-MM-DD\n"
        "fname,first_name,Direct\n"
        "salary,annual,salary × 12\n"
        "status,status,\"A → ACTIVE, I → INACTIVE, case insensitive\"\n"
    ).encode()
    typed = compile_rule_workbook(
        "types.csv",
        csv,
        source_columns=["dob", "fname", "salary", "status"],
        dest_columns=["age", "first_name", "annual", "status"],
        dest_types={"age": "INTEGER", "first_name": "TEXT", "annual": "NUMERIC", "status": "TEXT"},
        source_types={"dob": "DATE", "fname": "TEXT", "salary": "NUMERIC", "status": "TEXT"},
    )
    by = {r["source_column"]: r for r in typed["rules"] if r.get("source_column")}
    assert by["dob"]["status"] == "needs_confirmation"
    assert any("temporal" in issue for issue in by["dob"]["issues"])
    assert by["fname"]["status"] == "executable"
    assert by["salary"]["status"] == "executable"
    assert by["status"]["status"] == "needs_confirmation"
    assert any("case-insensitive" in issue.lower() or "exactly" in issue.lower() for issue in by["status"]["issues"])
    assert by["status"]["code_crosswalk"] == {"A": "ACTIVE", "I": "INACTIVE"}

    cdc = compile_rule_workbook(
        "cdc.csv",
        (
            "Source Column,Destination Column,Rule\n"
            "salary,annual,salary × 12\n"
            "fname,first_name,Direct\n"
        ).encode(),
        source_columns=["salary", "fname"],
        dest_columns=["annual", "first_name"],
        sync_mode="cdc",
    )
    by = {r["source_column"]: r for r in cdc["rules"] if r.get("source_column")}
    assert by["salary"]["status"] == "needs_confirmation"
    assert by["salary"]["shape_step"] is None
    assert any("cdc" in issue.lower() for issue in by["salary"]["issues"])
    assert by["fname"]["status"] == "executable"
    assert cdc["shape_steps"] == []
    assert cdc["sync_mode"] == "cdc"
    assert cdc["matcher"] == "cupid-linguistic+instance+type+constraint"


def test_compound_predicate_is_fail_closed_not_a_half_filter():
    both = classify_rule("keep if status = A and amount > 0")
    assert both["kind"] == "filter"
    assert "status = \"A\"" in both["condition"]
    assert "amount > 0" in both["condition"]
    assert "and" in both["condition"]
    leftover = classify_rule("keep if status in (A, I) and use the legacy flag")
    assert leftover["kind"] == "unknown"
    assert leftover["plane"] == "review"
    mixed = classify_rule("keep if status = A and amount > 0 or flag = 1")
    assert mixed["kind"] == "unknown"
    grouped = classify_rule("keep if (status = A or status = I) and amount > 0")
    assert grouped["kind"] == "filter"
    assert "or" in grouped["condition"]
    assert "amount > 0" in grouped["condition"]
    starts = classify_rule("keep if name starts with Acme")
    assert starts["kind"] == "filter"
    assert 'starts_with(name, "Acme")' in starts["condition"]
    contains = classify_rule("exclude rows where email contains test")
    assert contains["kind"] == "filter"
    assert contains["keep"] is False
    assert "contains(email" in contains["condition"]
    empty = classify_rule("keep if notes is empty")
    assert empty["condition"] == "is_null(notes)"
    ilike = classify_rule("keep if name ilike 'Acme%'")
    assert 'starts_with(lower(name), "acme")' in ilike["condition"]
    unquoted = classify_rule("keep if name like Acme%")
    assert 'starts_with(name, "Acme")' in unquoted["condition"]
    assert classify_rule("keep distinct rows")["kind"] == "unknown"
    assert classify_rule("pivot country into columns")["kind"] == "unknown"
    assert classify_rule("Y/N flag")["kind"] == "unknown"
    blank = classify_rule("treat empty as null")
    assert blank["kind"] == "null_if"
    assert blank["values"] == [""]
    coal = classify_rule("COALESCE(status, flag, 'UNK')")
    assert coal["kind"] == "derive"
    assert "coalesce(" in coal["expression"]
    assert '"UNK"' in coal["expression"]
    two = classify_rule("NVL(status, 'X')")
    assert two["kind"] == "default"
    assert two["value"] == "X"
    zone = classify_rule("assume timezone America/New_York")
    assert zone["kind"] == "timezone"
    assert zone["zone"] == "America/New_York"
    assert classify_rule("assume timezone")["plane"] == "review"

    csv = (
        "Source Column,Destination Column,Rule\n"
        "status,status,\"keep if status = A and amount > 0\"\n"
        "created_at,created_at,assume timezone America/New_York\n"
        "code,code,COALESCE(code, alt_code, 'UNK')\n"
        "notes,notes,treat empty as null\n"
        "pay,pay,Direct\n"
        "name,name,concat first last\n"
    ).encode()
    report = compile_rule_workbook(
        "pred-constraint.csv",
        csv,
        source_columns=["status", "amount", "created_at", "code", "alt_code", "notes", "pay", "name"],
        dest_columns=["status", "created_at", "code", "notes", "pay", "name"],
        source_types={"pay": "DECIMAL(18,6)", "name": "VARCHAR(80)"},
        dest_types={"pay": "DECIMAL(10,2)", "name": "VARCHAR(12)", "created_at": "TIMESTAMP"},
    )
    by = {r["source_column"]: r for r in report["rules"] if r.get("source_column")}
    filt = next(s for s in report["shape_steps"] if s["op"] == "filter_rows")
    assert "status = \"A\"" in filt["options"]["condition"]
    assert "amount > 0" in filt["options"]["condition"]
    assert by["created_at"]["status"] == "executable"
    assert by["created_at"]["transform"] == "assume_timezone"
    assert by["created_at"]["timezone"] == "America/New_York"
    assert by["created_at"]["engine_transform"] == "assume_timezone:America/New_York"
    assert by["code"]["kind"] == "derive"
    assert by["notes"]["kind"] == "null_if"
    assert by["pay"]["status"] == "needs_confirmation"
    assert any("precision" in issue.lower() or "scale" in issue.lower() for issue in by["pay"]["issues"])
    assert by["name"]["status"] == "needs_confirmation"
    assert any("varchar" in issue.lower() or "length" in issue.lower() for issue in by["name"]["issues"])


def test_volatile_identity_wrappers_and_lookup_payload():
    today = classify_rule("default to today")
    assert today["kind"] == "unknown"
    assert today["plane"] == "review"
    assert classify_rule("constant: GETDATE()")["kind"] == "unknown"
    assert classify_rule("NOW()")["kind"] == "unknown"
    assert classify_rule("skip deleted rows")["kind"] == "unknown"
    assert classify_rule("omit")["kind"] == "omit"
    assert classify_rule("# do not apply this note")["kind"] == "unknown"
    assert classify_rule("excel serial date")["kind"] == "unknown"

    wrapped = classify_rule("keep if trim(status) = A")
    assert wrapped["kind"] == "filter"
    assert wrapped["condition"] == 'trim(status) = "A"'
    nested = classify_rule("keep if upper(trim(status)) = ACTIVE")
    assert "upper(trim(status))" in nested["condition"]
    distinct = classify_rule("keep if status is distinct from flag")
    assert distinct["kind"] == "filter"
    assert "is_null(status)" in distinct["condition"]
    extrema = classify_rule("GREATEST(a, b, 0)")
    assert extrema["kind"] == "derive"
    assert extrema["expression"].startswith("greatest(")
    part = classify_rule("split_part(name, ' ', 1)")
    assert part["kind"] == "derive"
    assert "split_part(name" in part["expression"]

    csv = (
        "Source Column,Destination Column,Rule\n"
        "status,status,\"keep if trim(status) = A\"\n"
        "id,id,Direct\n"
        "code,amount,\"A → ACTIVE, I → INACTIVE\"\n"
        "pad,qty,\"01 → 1, 02 → 2\"\n"
        "low,high,GREATEST(low, high)\n"
    ).encode()
    report = compile_rule_workbook(
        "imap-integrity.csv",
        csv,
        source_columns=["status", "id", "code", "pad", "low", "high"],
        dest_columns=["status", "id", "amount", "qty", "high"],
        dest_types={
            "id": "INTEGER GENERATED ALWAYS AS IDENTITY",
            "amount": "INTEGER",
            "qty": "NUMERIC",
            "status": "TEXT",
            "high": "NUMERIC",
        },
    )
    by = {r["source_column"]: r for r in report["rules"] if r.get("source_column")}
    filt = next(s for s in report["shape_steps"] if s["op"] == "filter_rows")
    assert "trim(status)" in filt["options"]["condition"]
    assert by["id"]["status"] == "needs_confirmation"
    assert any("identity" in issue.lower() or "generated" in issue.lower() for issue in by["id"]["issues"])
    assert by["code"]["status"] == "needs_confirmation"
    assert any("numeric" in issue.lower() for issue in by["code"]["issues"])
    assert by["pad"]["status"] == "needs_confirmation"
    assert any("leading zero" in issue.lower() for issue in by["pad"]["issues"])
    assert by["low"]["kind"] == "derive"
    assert by["low"]["status"] == "executable"


def test_to_char_crypto_lpad_and_extract_are_not_silent():
    assert classify_rule("md5(email)")["kind"] == "hash"
    assert classify_rule("sha256(phone)")["kind"] == "hash"
    assert classify_rule("crc32(email)")["kind"] == "hash"
    assert classify_rule("Lowercase + validate email")["kind"] == "email"
    assert classify_rule("TO_CHAR(amount, '999,999.00')")["kind"] == "unknown"
    assert classify_rule("TO_CHAR(amount)")["kind"] == "unknown"
    dated = classify_rule("TO_CHAR(dob, 'YYYY-MM-DD')")
    assert dated["kind"] == "date"
    assert "YYYY" in (dated.get("format") or "")
    assert classify_rule("extract year from date")["kind"] == "unknown"
    assert classify_rule("DATEADD(day, 1, dob)")["kind"] == "unknown"
    assert classify_rule("soundex(name)")["kind"] == "unknown"
    assert classify_rule("european number")["kind"] == "unknown"
    assert classify_rule("COLLATE Latin1_General_CI_AS")["kind"] == "unknown"
    lpad = classify_rule("LPAD(code, 5, '0')")
    assert lpad["kind"] == "pad"
    assert lpad["width"] == 5
    assert lpad["side"] == "left"
    assert lpad["fill"] == "0"
    rpad = classify_rule("RPAD(name, 10)")
    assert rpad["kind"] == "pad"
    assert rpad["side"] == "right"
    assert rpad["fill"] == " "
    spoken = classify_rule("lpad code to 5 with 0")
    assert spoken["kind"] == "pad"
    assert spoken["fill"] == "0"
    english = classify_rule("pad left to 5")
    assert english["kind"] == "pad"
    assert english.get("fill", " ") == " "

    csv = (
        "Source Column,Destination Column,Rule\n"
        "email,email_hash,md5(email)\n"
        "code,code,\"LPAD(code, 5, '0')\"\n"
        "amount,amount,\"TO_CHAR(amount, '999,999.00')\"\n"
        "dob,birth_year,extract year from date\n"
    ).encode()
    report = compile_rule_workbook(
        "informatica-pad.csv",
        csv,
        source_columns=["email", "code", "amount", "dob"],
        dest_columns=["email_hash", "code", "amount", "birth_year"],
    )
    by = {r["source_column"]: r for r in report["rules"] if r.get("source_column")}
    assert by["email"]["transform"] == "hash_pii"
    assert by["email"]["status"] == "executable"
    pad_step = next(s for s in report["shape_steps"] if s["op"] == "pad")
    assert pad_step["options"]["fill"] == "0"
    assert pad_step["options"]["width"] == 5
    assert by["amount"]["status"] == "needs_confirmation"
    assert any("number mask" in issue.lower() or "parse_date" in issue.lower() for issue in by["amount"]["issues"])
    assert by["dob"]["status"] == "needs_confirmation"
    assert by["dob"]["kind"] != "date" or by["dob"]["status"] == "needs_confirmation"
    assert any("extract" in issue.lower() or "dateadd" in issue.lower() for issue in by["dob"]["issues"])


def test_scd_timezone_isnull_jsonpath_and_regex_are_not_silent():
    """Spark AutoCDC / DataCoolie SCD2 + Oracle AT TIME ZONE / FROM_TZ.

    Effective-date labels are history columns, not parse_date. A bare
    ``col is not null`` is filter-vs-contract, not a Validate write.
    JSON_VALUE is a path extract, not parse_json of the blob.
    """
    for text in (
        "effective date",
        "end date",
        "is_current",
        "scd type 2",
        "when matched then update",
        "slowly changing dimension",
    ):
        got = classify_rule(text)
        assert got["kind"] == "unknown", text
        assert got["plane"] == "review", text
        assert "scd" in got["reason"].lower() or "history" in got["reason"].lower(), text

    at_zone = classify_rule("AT TIME ZONE 'UTC'")
    assert at_zone["kind"] == "timezone"
    assert at_zone["zone"] == "UTC"
    from_tz = classify_rule("FROM_TZ(ts, 'UTC')")
    assert from_tz["kind"] == "timezone"
    assert from_tz["zone"] == "UTC"
    convert = classify_rule("convert timezone UTC to America/New_York")
    assert convert["kind"] == "unknown"
    assert convert["plane"] == "review"
    assert classify_rule("NEW_TIME(ts, 'EST', 'PST')")["kind"] == "unknown"
    assert classify_rule("time of day")["kind"] == "time"

    bare = classify_rule("deleted_at is not null")
    assert bare["kind"] == "unknown"
    assert bare["plane"] == "review"
    assert "ambiguous" in bare["reason"].lower()
    keep = classify_rule("keep if deleted_at is not null")
    assert keep["kind"] == "filter"
    assert keep["condition"] == "is_not_null(deleted_at)"
    assert classify_rule("NOT NULL")["kind"] == "contract"
    assert classify_rule("required")["kind"] == "contract"

    path = classify_rule("JSON_VALUE(payload, '$.city')")
    assert path["kind"] == "unknown"
    assert "json path" in path["reason"].lower() or "json_value" in path["reason"].lower()
    assert classify_rule("parse JSON")["kind"] == "json"
    assert classify_rule("xpath(xml, '//city')")["kind"] == "unknown"

    regex = classify_rule("keep if regex_matches(code, '^[A-Z]')")
    assert regex["kind"] == "filter"
    assert regex["condition"] == 'regex_matches(code, "^[A-Z]")'
    rlike = classify_rule("keep if code rlike 'A.*'")
    assert rlike["kind"] == "filter"
    assert "regex_matches(code" in rlike["condition"]
    tilde = classify_rule("keep if code ~ '^[A-Z]'")
    assert tilde["kind"] == "filter"

    assert classify_rule("row_number()")["kind"] == "unknown"
    assert classify_rule("lead(amount)")["kind"] == "unknown"
    assert classify_rule("listagg(name, ',')")["kind"] == "unknown"
    assert classify_rule("sequence.nextval")["kind"] == "unknown"
    assert classify_rule("base64 decode")["kind"] == "binary"
    assert classify_rule("DECODE(status, 'A', 'ACTIVE', 'I', 'INACTIVE')")["kind"] == "lookup"

    csv = (
        "Source Column,Destination Column,Rule\n"
        "created_at,created_at,AT TIME ZONE 'UTC'\n"
        "valid_from,valid_from,effective date\n"
        "deleted_at,deleted_at,deleted_at is not null\n"
        "payload,city,\"JSON_VALUE(payload, '$.city')\"\n"
        "code,code,\"keep if code ~ '^[A-Z]'\"\n"
        "blob,blob,base64 decode\n"
    ).encode()
    report = compile_rule_workbook(
        "scd-timezone-regex.csv",
        csv,
        source_columns=["created_at", "valid_from", "deleted_at", "payload", "code", "blob"],
        dest_columns=["created_at", "valid_from", "deleted_at", "city", "code", "blob"],
    )
    by = {r["source_column"]: r for r in report["rules"] if r.get("source_column")}
    assert by["created_at"]["transform"] == "assume_timezone"
    assert by["created_at"]["timezone"] == "UTC"
    assert by["created_at"]["status"] == "executable"
    assert by["valid_from"]["status"] == "needs_confirmation"
    assert by["valid_from"]["kind"] != "date"
    assert by["deleted_at"]["status"] == "needs_confirmation"
    assert by["deleted_at"]["kind"] != "contract"
    assert by["payload"]["status"] == "needs_confirmation"
    assert by["payload"]["kind"] != "json"
    assert by["blob"]["kind"] == "binary"
    assert by["blob"]["status"] == "executable"
    filt = next(s for s in report["shape_steps"] if s["op"] == "filter_rows")
    assert "regex_matches(code" in filt["options"]["condition"]


def test_na_utf8_tonumber_trycast_and_leftover_are_not_silent():
    """N/A ≠ passthrough; utf-8 ≠ utf minus 8; TRY_CAST swallows rows."""
    na = classify_rule("N/A")
    assert na["kind"] == "unknown"
    assert na["plane"] == "review"
    assert classify_rule("n/a")["kind"] == "unknown"
    assert classify_rule("not applicable")["kind"] == "unknown"
    assert classify_rule("Direct")["kind"] == "direct"
    assert classify_rule("same as source")["kind"] == "direct"
    assert classify_rule("copy from source")["kind"] == "direct"
    assert classify_rule("-")["kind"] == "direct"

    utf = classify_rule("convert to utf-8")
    assert utf["kind"] == "unknown"
    assert utf["plane"] == "review"
    assert "utf" in utf["reason"].lower() or "charset" in utf["reason"].lower()
    assert classify_rule("latin1 to utf8")["kind"] == "unknown"
    salary = classify_rule("salary × 12")
    assert salary["kind"] == "derive"
    assert "*" in salary["expression"]
    assert classify_rule("salary * 12")["kind"] == "derive"

    leftover = classify_rule("nvl(status,'X') leftover extra")
    assert leftover["kind"] == "unknown"
    assert leftover["plane"] == "review"
    assert classify_rule("NVL(status, 'X')")["kind"] == "default"
    coal = classify_rule("coalesce empty")
    assert coal["kind"] == "unknown"

    numbered = classify_rule("to_number(amount, '999D99')")
    assert numbered["kind"] == "unknown"
    assert numbered["plane"] == "review"
    assert classify_rule("to_number(amount)")["kind"] == "cast_number"
    assert classify_rule("VALUE(amount)")["kind"] == "cast_number"

    wrapped = classify_rule("nullif(trim(status), '')")
    assert wrapped["kind"] == "null_if"
    assert wrapped["values"] == [""]
    extras = {item.get("kind") or item.get("op") for item in wrapped.get("extras") or []}
    assert "trim" in extras

    quoted = classify_rule("iif(isnull(status), 'X', status)")
    assert quoted["kind"] == "derive"
    assert '"X"' in quoted["expression"]
    iff = classify_rule("iff(status='A','Y','N')")
    assert iff["kind"] == "derive"
    assert '"Y"' in iff["expression"] or "Y" in iff["expression"]

    assert classify_rule("try_cast(amount as number)")["kind"] == "unknown"
    assert classify_rule("safe_cast(amount as number)")["kind"] == "unknown"
    assert classify_rule("generate uuid")["kind"] == "unknown"
    assert classify_rule("btrim(name)")["kind"] == "trim"

    csv = (
        "Source Column,Destination Column,Rule\n"
        "notes,notes,N/A\n"
        "name,name,convert to utf-8\n"
        "amount,amount,\"to_number(amount, '999D99')\"\n"
        "status,status,\"nvl(status,'X') leftover extra\"\n"
        "flag,flag,\"iif(isnull(flag), 'X', flag)\"\n"
        "qty,qty,try_cast(qty as number)\n"
    ).encode()
    report = compile_rule_workbook(
        "na-utf8-trycast.csv",
        csv,
        source_columns=["notes", "name", "amount", "status", "flag", "qty"],
        dest_columns=["notes", "name", "amount", "status", "flag", "qty"],
    )
    by = {r["source_column"]: r for r in report["rules"] if r.get("source_column")}
    assert by["notes"]["status"] == "needs_confirmation"
    assert by["notes"]["kind"] != "direct"
    assert by["name"]["status"] == "needs_confirmation"
    assert by["name"]["kind"] != "derive"
    assert by["amount"]["status"] == "needs_confirmation"
    assert by["amount"]["kind"] != "cast_number"
    assert by["status"]["status"] == "needs_confirmation"
    assert by["flag"]["kind"] == "derive"
    assert by["flag"]["status"] == "executable"
    assert by["qty"]["status"] == "needs_confirmation"


def test_multi_table_catalog_bind_is_fail_closed_on_homonyms():
    """Clio object+attribute: id on two tables is not a unique winner."""
    csv = (
        "Source,Source Column,Destination,Destination Column,Rule\n"
        "customers,id,dw,customer_id,Direct\n"
        "orders,id,dw,order_id,Direct\n"
        ",id,dw,mystery_id,Direct\n"
        "ghost,name,dw,name,Direct\n"
        "customers,email,dw,email,Direct\n"
    ).encode()
    catalog = {
        "customers": ["id", "email", "name"],
        "orders": ["id", "amount", "customer_id"],
    }
    report = compile_rule_workbook(
        "multi-table.csv",
        csv,
        source_columns=["id", "email", "name", "amount", "customer_id"],
        dest_columns=["customer_id", "order_id", "mystery_id", "name", "email"],
        source_tables=["customers", "orders"],
        source_catalog=catalog,
        dest_table="dw",
    )
    by = {(r.get("source_table"), r.get("source_column")): r for r in report["rules"]}
    assert by[("customers", "id")]["status"] == "executable"
    assert by[("customers", "id")]["dest_column"] == "customer_id"
    assert by[("orders", "id")]["status"] == "executable"
    assert by[("orders", "id")]["dest_column"] == "order_id"
    mystery = next(r for r in report["rules"] if r.get("dest_column") == "mystery_id" or "mystery" in (r.get("rule_text") or "").lower() or (not r.get("source_table") and r.get("source_column") == "id"))
    assert mystery["status"] == "needs_confirmation"
    assert any("customers" in i and "orders" in i for i in mystery["issues"])
    ghost = next(r for r in report["rules"] if r.get("source_table") == "ghost")
    assert ghost["status"] == "needs_confirmation"
    assert by[("customers", "email")]["status"] == "executable"
    assert report["source_tables"] == ["customers", "orders"]


def test_comma_source_table_is_parsed_as_selected_tables():
    csv = (
        "Source,Source Column,Destination Column,Rule\n"
        ",email,email,Direct\n"
    ).encode()
    report = compile_rule_workbook(
        "comma-tables.csv",
        csv,
        source_columns=["id", "email"],
        dest_columns=["email"],
        source_table="customers, orders",
        source_catalog={"customers": ["id", "email"], "orders": ["id", "amount"]},
    )
    email = next(r for r in report["rules"] if r.get("source_column") == "email")
    assert email["source_table"] == "customers"
    assert email["status"] == "executable"


def test_dest_catalog_and_table_scoped_edges_are_not_silent():
    csv = (
        "Source,Source Column,Destination,Destination Column,Rule\n"
        "customers,id,dim_customer,id,Direct\n"
        "orders,id,fact_order,id,Direct\n"
        "customers,pay,dim_customer,amount,Direct\n"
        ",ghost_col,dim_customer,missing,Direct\n"
    ).encode()
    report = compile_rule_workbook(
        "dest-catalog.csv",
        csv,
        source_tables=["customers", "orders"],
        source_catalog={"customers": ["id", "pay"], "orders": ["id", "amount"]},
        dest_tables=["dim_customer", "fact_order"],
        dest_catalog={
            "dim_customer": ["id", "email", "amount"],
            "fact_order": ["id", "customer_id", "amount"],
        },
        dest_types={
            "dim_customer.amount": "NUMERIC(10,2)",
            "customers.pay": "NUMERIC(18,6)",
        },
        source_types={
            "customers.pay": "NUMERIC(18,6)",
            "dim_customer.amount": "NUMERIC(10,2)",
        },
    )
    by = {
        (r.get("source_table"), r.get("source_column"), r.get("dest_table")): r
        for r in report["rules"]
    }
    assert by[("customers", "id", "dim_customer")]["status"] == "executable"
    assert by[("orders", "id", "fact_order")]["status"] == "executable"
    assert by[("customers", "id", "dim_customer")]["dest_column"] == "id"
    assert by[("orders", "id", "fact_order")]["dest_column"] == "id"
    pay = by[("customers", "pay", "dim_customer")]
    assert pay["status"] == "needs_confirmation"
    assert any("precision" in i.lower() or "scale" in i.lower() for i in pay["issues"])
    ghost = next(r for r in report["rules"] if r.get("source_column") == "ghost_col")
    assert ghost["status"] == "needs_confirmation"
    tables = {item["source_table"] for item in report["projection"]}
    assert "customers" in tables and "orders" in tables


def test_shape_steps_are_stamped_with_source_table():
    csv = (
        "Source,Source Column,Destination Column,Rule\n"
        "customers,signed_on,birth_date,MM/DD/YYYY\n"
        "customers,email,email,lowercase\n"
        "orders,amount,amount,Direct\n"
    ).encode()
    report = compile_rule_workbook(
        "shape-scope.csv",
        csv,
        source_tables=["customers", "orders"],
        source_catalog={
            "customers": ["id", "email", "signed_on"],
            "orders": ["id", "amount"],
        },
        dest_columns=["birth_date", "email", "amount"],
    )
    date_step = next(s for s in report["shape_steps"] if s["op"] == "parse_date")
    assert date_step["source_table"] == "customers"
    assert date_step["column"] == "signed_on"
    by_src = {r["source_column"]: r for r in report["rules"] if r.get("source_table") == "customers"}
    assert by_src["signed_on"]["shape_step"]["source_table"] == "customers"
    grouped = {item["source_table"]: item["steps"] for item in report["shape_steps_by_table"]}
    assert "customers" in grouped
    assert all(s.get("source_table") == "customers" for s in grouped["customers"])
    assert "orders" not in grouped or not any(
        s.get("op") == "parse_date" for s in grouped.get("orders", [])
    )


def test_apply_projection_is_table_scoped_and_fail_closed():
    from services.rule_compiler.apply import apply_compiled_projection, apply_selected_tables

    csv = (
        "Source,Source Column,Destination,Destination Column,Rule\n"
        "customers,id,dim_customer,customer_id,Direct\n"
        "customers,email,dim_customer,email,lowercase\n"
        "customers,signed_on,dim_customer,birth_date,MM/DD/YYYY\n"
        "customers,status,dim_customer,status,A → ACTIVE, I → INACTIVE\n"
        "orders,id,fact_order,order_id,Direct\n"
        "orders,amount,fact_order,amount,Direct\n"
        "orders,status,fact_order,status,A → OPEN, I → CLOSED\n"
        ",id,dw,mystery_id,Direct\n"
    ).encode()
    report = compile_rule_workbook(
        "apply-scope.csv",
        csv,
        source_tables=["customers", "orders"],
        source_catalog={
            "customers": ["id", "email", "signed_on", "status"],
            "orders": ["id", "amount", "status"],
        },
        dest_tables=["dim_customer", "fact_order"],
        dest_catalog={
            "dim_customer": ["customer_id", "email", "birth_date", "status"],
            "fact_order": ["order_id", "amount", "status"],
        },
    )
    mystery = next(r for r in report["rules"] if r.get("dest_column") == "mystery_id")
    assert mystery["status"] == "needs_confirmation"

    customers = [
        {"id": 1, "email": "Ada@Example.COM", "signed_on": "03/15/2020", "status": "A"},
        {"id": 2, "email": "bad@example.com", "signed_on": "not-a-date", "status": "A"},
        {"id": 3, "email": "c@example.com", "signed_on": "04/01/2021", "status": "Z"},
    ]
    orders = [
        {"id": 10, "amount": "19.50", "status": "A"},
        {"id": 11, "amount": "4.00", "status": "Z"},
    ]
    applied = apply_selected_tables(
        report,
        {"customers": customers, "orders": orders},
    )
    assert applied["missing_populations"] == []
    by_src = {item["source_table"]: item for item in applied["tables"]}

    cust_dest = by_src["customers"]["destinations"][0]
    assert cust_dest["dest_table"] == "dim_customer"
    assert len(cust_dest["rows"]) == 1
    row = cust_dest["rows"][0]
    assert row["customer_id"] == 1
    assert row["email"] == "ada@example.com"
    assert row["birth_date"] == "2020-03-15"
    assert row["status"] == "ACTIVE"
    assert "mystery_id" not in row
    assert cust_dest["refused"] >= 2
    assert any("not-a-date" in (q.get("error") or "") or q.get("column") == "birth_date" for q in cust_dest["quarantine"])
    assert any(q.get("column") == "status" and "unmapped" in (q.get("error") or "") for q in cust_dest["quarantine"])

    ord_dest = by_src["orders"]["destinations"][0]
    assert ord_dest["dest_table"] == "fact_order"
    assert len(ord_dest["rows"]) == 1
    assert ord_dest["rows"][0]["order_id"] == 10
    assert ord_dest["rows"][0]["status"] == "OPEN"
    assert "email" not in ord_dest["rows"][0]
    assert "birth_date" not in ord_dest["rows"][0]
    assert any(q.get("column") == "status" for q in ord_dest["quarantine"])

    leaked = apply_compiled_projection(report, source_table="orders", rows=customers)
    leaked_row = leaked["destinations"][0]["rows"][0] if leaked["destinations"][0]["rows"] else {}
    assert "email" not in leaked_row
    assert "birth_date" not in leaked_row
